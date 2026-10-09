"""Génère des préconisations en français à partir des faits déjà calculés
par Checkers.py (PSet manquants) et carbon.py (bilan carbone, scénarios de
substitution chiffrés), avec le même schéma local-first / cloud plafonné
validé dans ~/Desktop/mon-projet. Trois agents "experts" (carbone, qualité
des données, métier/risques projet) se prononcent chacun sur une facette
distincte des mêmes faits, puis un agent de synthèse croise leurs constats.
Exécutés séquentiellement, pas en parallèle : mesuré empiriquement que ça
n'apporte rien avec Ollama (serveur local qui traite les requêtes une par
une de toute façon) et ralentit à cause de la charge supplémentaire.
Aucun agent ne recalcule ni n'invente de chiffre -- toute la logique
numérique (scénarios de substitution, priorités métier) est calculée en
Python déterministe dans ce module et dans carbon.py. Pour chaque expert,
les chiffres eux-mêmes sont rédigés en liste à puces par Python
(_format_*_bullets) et jamais reformulés par le LLM ; l'agent ne rédige
qu'une courte intro (1-2 phrases) de mise en perspective, concaténée devant
la liste. Ça garantit des chiffres toujours exacts quel que soit le palier
(local/cloud), et laisse au petit modèle local une tâche plus simple qu'une
liste complète -- constaté empiriquement que lui demander de transcrire ET
la liste ET une mise en garde produisait des réponses vagues ou déformées.

Toute dépendance à `crewai` est importée paresseusement (à l'intérieur des
fonctions), pas au niveau du module : si ENABLE_AI_RECOMMENDATIONS=false,
l'app doit pouvoir démarrer et fonctionner même si `crewai` n'est pas
installé.
"""

import os

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from carbon import MATERIAL_CARBON_FACTORS, compute_substitution_scenarios
from excel_safety import excel_safe

_KEEP_NAMES_INSTRUCTION = (
    "Réponds en français. Garde les noms d'étages, de types d'éléments, de matériaux et de "
    "PSet exactement tels qu'ils apparaissent ci-dessous (ne les traduis pas, ne les modifie pas)."
)


def _ai_enabled() -> bool:
    return os.environ.get("ENABLE_AI_RECOMMENDATIONS", "true").strip().lower() not in ("false", "0", "no")


def _strip_bullet_lines(text: str) -> str:
    """Retire les lignes à puces d'une intro rédigée par un agent. Constaté
    empiriquement : même invité explicitement à ne pas le faire, un petit
    modèle local recopie parfois la liste qu'on lui a donnée en contexte --
    on ne peut pas compter sur la consigne seule, donc on nettoie en Python
    plutôt que de risquer un doublon visible dans le rapport."""
    return "\n".join(
        line for line in text.splitlines()
        if line.strip() and not line.strip().startswith(("-", "*"))
    ).strip()


def compute_business_priorities(elements_by_class: dict, carbon_by_type: dict,
                                 validity_threshold: float = 0.8, top_n: int = 5) -> list:
    """Croise deux faits déjà calculés (taux de validité par type, empreinte
    carbone par type) pour identifier ce qui compte le plus sur le terrain :
    les types d'éléments à la fois mal renseignés ET carbonés. Fonction pure
    (pas d'appel LLM), testable isolément."""
    priorities = []
    for ifc_class, stats in elements_by_class.items():
        total = stats.get("total", 0)
        if total == 0:
            continue
        valid_rate = stats.get("valid", 0) / total
        carbon_kg = carbon_by_type.get(ifc_class, 0)
        if valid_rate < validity_threshold and carbon_kg != 0:
            priorities.append({
                "ifc_class": ifc_class,
                "valid_rate": valid_rate,
                "carbon_kg": carbon_kg,
                "total_elements": total,
            })
    priorities.sort(key=lambda p: abs(p["carbon_kg"]), reverse=True)
    return priorities[:top_n]


def _build_vocabulary(validation_facts: dict, carbon_facts: dict) -> set:
    vocab = set()
    vocab |= {str(c).lower() for c in carbon_facts.get("by_type", {})}
    vocab |= {str(m).lower() for m in carbon_facts.get("by_material", {})}
    vocab |= {str(f).lower() for f in carbon_facts.get("by_floor", {})}
    vocab |= {m.lower() for m in MATERIAL_CARBON_FACTORS}
    vocab |= {str(c).lower() for c in validation_facts.get("elements_by_class", {})}
    for item in validation_facts.get("top_missing", []):
        vocab.add(str(item["floor"]).lower())
        vocab.add(str(item["ifc_class"]).lower())
        vocab.add(str(item["pset"]).lower())
    return vocab


def _format_carbon_bullets(scenarios: list) -> str:
    """Lignes 100% déterministes (Python) -- jamais passées par un LLM, donc
    jamais de risque de chiffre déformé. C'est ce qui s'affiche tel quel dans
    le rapport, sous l'intro rédigée par l'agent (voir _carbon_expert_step)."""
    if not scenarios:
        return "- Aucun scénario de substitution calculable pour ce projet."
    return "\n".join(
        f"- {s['from_material']} -> {s['to_material']} : {s['current_kg']:.0f} -> "
        f"{s['alternative_kg']:.0f} kg CO2e ({s['delta_pct']:+.0f} %)"
        for s in scenarios
    )


def _carbon_expert_step(carbon_facts: dict, scenarios: list):
    from crewai import Agent, Task

    top_scenarios = scenarios[:3]
    bullets = _format_carbon_bullets(top_scenarios)

    def build(llm):
        agent = Agent(
            role="Expert carbone BIM",
            goal="Mettre en perspective des scénarios de réduction carbone déjà chiffrés et priorisés.",
            backstory=(
                "Un ingénieur carbone qui travaille uniquement à partir des chiffres déjà calculés "
                "qu'on lui fournit -- il n'invente ni ne recalcule jamais un chiffre."
            ),
            llm=llm,
            verbose=False,
        )
        task = Task(
            description=(
                "Voici des scénarios de substitution de matériaux déjà chiffrés et déjà rédigés en liste "
                f"à puces (ce sont des comparaisons d'ordre de grandeur, pas des préconisations "
                f"d'ingénierie prêtes à l'emploi) :\n\n{bullets}\n\n{_KEEP_NAMES_INSTRUCTION}\n"
                "Rédige UNIQUEMENT 1 à 2 phrases d'introduction qui indiquent lequel de ces scénarios est "
                "le plus prioritaire et pourquoi. Ne réécris PAS la liste ci-dessus, elle sera affichée "
                "telle quelle juste après tes phrases. Ne cite aucun autre chiffre que ceux ci-dessus."
            ),
            expected_output="1 à 2 phrases en français, sans liste à puces, sans chiffre inventé.",
            agent=agent,
        )
        return agent, task

    must_include = [top_scenarios[0]["from_material"]] if top_scenarios else []
    return build, bullets, must_include


def _format_quality_bullets(top_missing: list, reliability: dict) -> str:
    lines = [
        f"- Étage {item['floor']}, {item['ifc_class']} : PSet '{item['pset']}' manquant ou incomplet "
        f"sur {item['count']} élément(s)"
        for item in top_missing
    ] or ["- Aucune non-conformité de PSet trouvée."]
    lines.append(
        f"- Fiabilité du bilan carbone : {reliability.get('pct_mesure', 0):.0f} % mesuré, "
        f"{reliability.get('pct_estime', 0):.0f} % estimé, "
        f"{reliability.get('pct_non_calculable', 0):.0f} % non calculable"
    )
    return "\n".join(lines)


def _data_quality_expert_step(validation_facts: dict, carbon_facts: dict):
    from crewai import Agent, Task

    top_missing = validation_facts.get("top_missing", [])[:8]
    reliability = carbon_facts.get("reliability_summary", {})
    bullets = _format_quality_bullets(top_missing, reliability)

    def build(llm):
        agent = Agent(
            role="Expert qualité des données BIM",
            goal="Mettre en perspective les corrections de maquette et la fiabilité des données déjà calculées.",
            backstory=(
                "Un coordinateur BIM qui priorise exactement les manques déjà détectés qu'on lui donne "
                "-- il n'en ajoute ni n'en retire selon son propre jugement."
            ),
            llm=llm,
            verbose=False,
        )
        task = Task(
            description=(
                "Voici le décompte déjà vérifié des PSet/paramètres manquants et la fiabilité du bilan "
                f"carbone, déjà rédigés en liste à puces :\n\n{bullets}\n\n{_KEEP_NAMES_INSTRUCTION}\n"
                "Rédige UNIQUEMENT 1 à 2 phrases d'introduction indiquant quelle correction est la plus "
                "urgente et pourquoi. Ne réécris PAS la liste ci-dessus, elle sera affichée telle quelle "
                "juste après tes phrases. Ne cite aucun autre chiffre que ceux ci-dessus."
            ),
            expected_output="1 à 2 phrases en français, sans liste à puces, sans chiffre inventé.",
            agent=agent,
        )
        return agent, task

    must_include = [str(top_missing[0]["ifc_class"])] if top_missing else []
    return build, bullets, must_include


def _format_business_bullets(business_priorities: list) -> str:
    if not business_priorities:
        return "- Aucun type d'élément ne cumule à la fois un faible taux de conformité et un impact carbone notable."
    return "\n".join(
        f"- {p['ifc_class']} : {p['valid_rate']*100:.0f} % conforme sur {p['total_elements']} éléments, "
        f"{p['carbon_kg']:.0f} kg CO2e d'empreinte carbone"
        for p in business_priorities
    )


def _business_expert_step(business_priorities: list):
    from crewai import Agent, Task

    bullets = _format_business_bullets(business_priorities)

    def build(llm):
        agent = Agent(
            role="Expert métier / risques projet BIM",
            goal="Mettre en perspective ce qui mérite une attention immédiate, en croisant qualité de maquette et impact carbone.",
            backstory=(
                "Un chef de projet BIM qui priorise le travail de l'équipe à partir de faits déjà croisés "
                "-- il n'évalue pas la faisabilité technique des solutions, seulement l'urgence à traiter."
            ),
            llm=llm,
            verbose=False,
        )
        task = Task(
            description=(
                "Voici les types d'éléments déjà identifiés comme cumulant un faible taux de conformité "
                f"ET un impact carbone notable, déjà rédigés en liste à puces :\n\n{bullets}\n\n"
                f"{_KEEP_NAMES_INSTRUCTION}\n"
                "Rédige UNIQUEMENT 1 à 2 phrases d'introduction expliquant pourquoi le premier type de la "
                "liste mérite une attention immédiate (ex. données manquantes sur un poste à fort impact "
                "carbone -- risque de mauvaise estimation, blocage du métré/de la facturation). Ne réécris "
                "PAS la liste ci-dessus, elle sera affichée telle quelle juste après tes phrases."
            ),
            expected_output="1 à 2 phrases en français, sans liste à puces, sans chiffre inventé.",
            agent=agent,
        )
        return agent, task

    must_include = [business_priorities[0]["ifc_class"]] if business_priorities else []
    return build, bullets, must_include


def _synthesis_step(carbon_text: str, quality_text: str, business_text: str):
    from crewai import Agent, Task

    def build(llm):
        agent = Agent(
            role="Rédacteur de synthèse",
            goal="Croiser les avis de trois experts en une note priorisée pour un chef de projet BIM.",
            backstory=(
                "Un rédacteur technique qui ne se contente pas de juxtaposer les avis de plusieurs "
                "spécialistes : il fait ressortir où ils se recoupent et où ils créent une tension."
            ),
            llm=llm,
            verbose=False,
        )
        task = Task(
            description=(
                "Voici les avis de trois experts sur le même projet. Rédige une note Markdown en français "
                "pour un chef de projet BIM. N'ajoute aucun constat qui n'est pas mentionné par l'un des "
                "trois avis ci-dessous.\n\n"
                f"Avis de l'expert carbone :\n{carbon_text}\n\n"
                f"Avis de l'expert qualité des données :\n{quality_text}\n\n"
                f"Avis de l'expert métier / risques projet :\n{business_text}\n\n"
                "Structure la note ainsi : une section '## Priorités croisées' (2-4 points où plusieurs "
                "experts se recoupent ou se contredisent -- explique la tension si elle existe), puis "
                "'## Avis carbone', '## Avis qualité des données' et '## Avis métier' reprenant les points "
                "clés de chaque avis."
            ),
            expected_output="Une note Markdown en français avec les quatre sections décrites.",
            agent=agent,
        )
        return agent, task

    headers = ["Priorités croisées", "Avis carbone", "Avis qualité des données", "Avis métier"]
    return build, headers


def run_recommendations_pipeline(validation_facts: dict, carbon_facts: dict, elements_by_class: dict) -> dict:
    """Point d'entrée. Ne lève jamais d'exception -- en cas de problème
    (agents désactivés, Ollama/cloud indisponibles, etc.), retourne un
    résultat "degraded" avec la raison, pour que /upload continue de
    fonctionner normalement (validation PSet + bilan carbone) sans les
    préconisations."""
    result = {
        "carbon_expert": None,
        "data_quality_expert": None,
        "business_expert": None,
        "synthesis": None,
        "tiers_used": {},
        "degraded": True,
        "degraded_reason": None,
    }

    if not _ai_enabled():
        result["degraded_reason"] = "Préconisations IA désactivées (ENABLE_AI_RECOMMENDATIONS=false)"
        return result

    try:
        from .llm_tiers import run_with_fallback
        from .validators import grounded, has_sections
    except Exception as exc:
        result["degraded_reason"] = f"Dépendances agents indisponibles ({exc})"
        return result

    os.environ.setdefault("OPENAI_API_KEY", "NA")  # crewai/litellm en ont besoin même sans l'utiliser

    scenarios = compute_substitution_scenarios(carbon_facts)
    business_priorities = compute_business_priorities(elements_by_class, carbon_facts.get("by_type", {}))
    vocabulary = _build_vocabulary(validation_facts, carbon_facts)
    tiers = {}

    # Exécution séquentielle : les 3 experts sont indépendants et pourraient
    # tourner en parallèle, mais mesuré empiriquement sur ce projet avec
    # Ollama, ça ralentit (le serveur local traite les requêtes une par une
    # de toute façon) sans aucun gain -- donc pas de ThreadPoolExecutor ici.
    # Chaque expert = une courte intro rédigée par l'agent + une liste à
    # puces déjà rédigée par Python (voir _format_*_bullets), jamais
    # reformulée par le LLM. Ça garantit que les chiffres affichés sont
    # toujours exacts, quel que soit le palier (local/cloud) ou le résultat
    # du grounding sur l'intro -- et ça laisse au petit modèle local une
    # tâche plus simple (1-2 phrases) qu'il gère mieux qu'une liste complète.
    try:
        build, carbon_bullets, must_include = _carbon_expert_step(carbon_facts, scenarios)
        carbon_intro, tiers["carbon"] = run_with_fallback(
            build, lambda text: grounded(text, must_include=must_include, vocabulary=vocabulary)
        )
        carbon_text = f"{_strip_bullet_lines(carbon_intro)}\n\n{carbon_bullets}"

        build, quality_bullets, must_include = _data_quality_expert_step(validation_facts, carbon_facts)
        quality_intro, tiers["data_quality"] = run_with_fallback(
            build, lambda text: grounded(text, must_include=must_include, vocabulary=vocabulary)
        )
        quality_text = f"{_strip_bullet_lines(quality_intro)}\n\n{quality_bullets}"

        build, business_bullets, must_include = _business_expert_step(business_priorities)
        business_intro, tiers["business"] = run_with_fallback(
            build, lambda text: grounded(text, must_include=must_include, vocabulary=vocabulary)
        )
        business_text = f"{_strip_bullet_lines(business_intro)}\n\n{business_bullets}"

        build, headers = _synthesis_step(carbon_text, quality_text, business_text)
        synthesis_text, tiers["synthesis"] = run_with_fallback(build, lambda text: has_sections(text, headers))
    except Exception as exc:
        result["degraded_reason"] = f"Pipeline de préconisations indisponible ({exc})"
        result["tiers_used"] = tiers
        return result

    return {
        "carbon_expert": carbon_text,
        "data_quality_expert": quality_text,
        "business_expert": business_text,
        "synthesis": synthesis_text,
        "tiers_used": tiers,
        "degraded": False,
        "degraded_reason": None,
    }


def _write_bullets(sheet, start_row: int, title: str, text, header_fill, header_font, normal_font, border) -> int:
    row = start_row
    cell = sheet.cell(row=row, column=1)
    cell.value = title
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1

    if not text:
        cell = sheet.cell(row=row, column=1)
        cell.value = "Non disponible."
        cell.font = normal_font
        cell.border = border
        row += 1
        return row + 1

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        cell = sheet.cell(row=row, column=1)
        cell.value = excel_safe(line.lstrip("-#* ").strip() if line.lstrip().startswith(("-", "*")) else line)
        cell.font = normal_font
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1

    return row + 1


def create_recommendations_sheet(workbook, recommendations_data: dict):
    """Nouvel onglet Excel "Préconisations" : synthèse croisée puis les trois
    avis d'experts (carbone, qualité des données, métier), et un bandeau de
    transparence sur le palier LLM utilisé (local/cloud) par étape."""
    sheet = workbook.create_sheet("Préconisations")

    header_font = Font(bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")
    normal_font = Font(size=11)
    footer_font = Font(size=10, italic=True, color="555555")
    border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    for col in range(1, 5):
        sheet.column_dimensions[get_column_letter(col)].width = 30

    row = 1
    title_cell = sheet.cell(row=row, column=1)
    title_cell.value = "PRÉCONISATIONS"
    title_cell.font = header_font
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    sheet.row_dimensions[row].height = 30
    row += 2

    if recommendations_data.get("degraded"):
        cell = sheet.cell(row=row, column=1)
        cell.value = f"Préconisations IA non disponibles : {recommendations_data.get('degraded_reason') or 'raison inconnue'}"
        cell.font = footer_font
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        return

    row = _write_bullets(
        sheet, row, "Synthèse (priorités croisées)",
        recommendations_data.get("synthesis"),
        header_fill, header_font, normal_font, border,
    )
    row = _write_bullets(
        sheet, row, "Avis de l'expert carbone",
        recommendations_data.get("carbon_expert"),
        header_fill, header_font, normal_font, border,
    )
    row = _write_bullets(
        sheet, row, "Avis de l'expert qualité des données",
        recommendations_data.get("data_quality_expert"),
        header_fill, header_font, normal_font, border,
    )
    row = _write_bullets(
        sheet, row, "Avis de l'expert métier / risques projet",
        recommendations_data.get("business_expert"),
        header_fill, header_font, normal_font, border,
    )

    tiers = recommendations_data.get("tiers_used", {})
    tiers_text = ", ".join(f"{step}: {tier}" for step, tier in tiers.items())
    footer_cell = sheet.cell(row=row, column=1)
    footer_cell.value = f"Palier LLM utilisé par étape -- {tiers_text or 'inconnu'}"
    footer_cell.font = footer_font
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
