"""Contrôle d'un modèle IFC contre un fichier IDS (Information Delivery
Specification, standard buildingSMART) via `ifctester`, le validateur de
référence d'IfcOpenShell.

Un IDS exprime, de façon ouverte et lisible par machine, ce qu'un modèle doit
contenir (quelles classes, quels PSet/propriétés/attributs, avec quelles
valeurs). C'est l'équivalent normalisé de l'Excel de règles de Checkers :
les deux peuvent être utilisés seuls ou ensemble.

Le fichier IDS est traité comme une entrée **non fiable** (typiquement reçu
d'un maître d'ouvrage ou d'un tiers) : voir `load_ids` pour les garde-fous.
"""

import re
from collections import Counter
from typing import Callable, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_safety import excel_safe

# Un IDS réel fait quelques dizaines de Ko ; 10 Mo laisse une très large marge.
MAX_IDS_BYTES = 10 * 1024 * 1024
# Longueur maximale d'un motif regex (xs:pattern) accepté.
MAX_PATTERN_LENGTH = 256
# Plafond de lignes de détail (une par élément non conforme) conservées par
# lot pour l'onglet Excel -- les compteurs, eux, restent exacts.
MAX_DETAIL_ROWS_PER_LOT = 5000


class IdsFileError(ValueError):
    """Fichier IDS refusé (illisible, invalide, ou jugé dangereux)."""


_FORBIDDEN_XML_DECLARATIONS = ("<!doctype", "<!entity")
_XML_ENCODING_RE = re.compile(r"""<\?xml[^>]*\bencoding\s*=\s*["']([^"']+)["']""", re.I)
_PATTERN_VALUE_RE = re.compile(
    r"""<(?:[\w.-]+:)?pattern\b[^>]*?\bvalue\s*=\s*(?:"([^"]*)"|'([^']*)')""", re.I
)
# Groupe contenant un quantificateur, lui-même quantifié : ((a+)+, (x*)*, ...).
# Heuristique anti-ReDoS : n'attrape pas tous les cas, mais écarte la forme
# classique du backtracking catastrophique.
_NESTED_QUANTIFIER_RE = re.compile(r"\((?:[^()\\]|\\.)*[+*](?:[^()\\]|\\.)*\)\s*(?:[+*]|\{\d+,\d*\})")


def _check_raw_ids_text(raw: bytes) -> str:
    if len(raw) > MAX_IDS_BYTES:
        raise IdsFileError(f"Fichier IDS trop volumineux (max {MAX_IDS_BYTES // (1024 * 1024)} Mo)")
    # Un XML UTF-8 ne contient jamais d'octet NUL : sa présence trahit un
    # encodage UTF-16/32, sous lequel la recherche de "<!DOCTYPE" ci-dessous
    # serait contournable (caractères entrelacés d'octets nuls).
    if b"\x00" in raw or raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        raise IdsFileError("Encodage du fichier IDS non supporté (UTF-8 requis)")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IdsFileError("Encodage du fichier IDS non supporté (UTF-8 requis)")

    declared = _XML_ENCODING_RE.search(text[:200])
    if declared and declared.group(1).lower().replace("_", "-") not in ("utf-8", "utf8", "us-ascii", "ascii"):
        raise IdsFileError("Encodage déclaré du fichier IDS non supporté (UTF-8 requis)")

    # Les IDS n'ont aucune raison de déclarer un DTD ou des entités ; les
    # refuser écarte d'emblée XXE et expansion d'entités ("billion laughs").
    lowered = text.lower()
    if any(token in lowered for token in _FORBIDDEN_XML_DECLARATIONS):
        raise IdsFileError("Fichier IDS refusé : déclaration DOCTYPE/ENTITY interdite")
    return text


def _check_patterns(text: str) -> None:
    for match in _PATTERN_VALUE_RE.finditer(text):
        pattern = match.group(1) if match.group(1) is not None else match.group(2)
        if len(pattern) > MAX_PATTERN_LENGTH:
            raise IdsFileError(f"Fichier IDS refusé : motif regex trop long (max {MAX_PATTERN_LENGTH} caractères)")
        if _NESTED_QUANTIFIER_RE.search(pattern):
            raise IdsFileError("Fichier IDS refusé : motif regex à quantificateurs imbriqués (risque de blocage)")


def load_ids(file_path: str):
    """Ouvre et valide (schéma XSD buildingSMART) un fichier IDS non fiable.
    Lève IdsFileError avec un message présentable à l'utilisateur."""
    from ifctester import ids

    with open(file_path, "rb") as fh:
        raw = fh.read(MAX_IDS_BYTES + 1)
    text = _check_raw_ids_text(raw)
    _check_patterns(text)

    try:
        document = ids.open(file_path, validate=True)
    except Exception as exc:
        message = str(exc).strip()
        reason = message.splitlines()[0] if message else type(exc).__name__
        # Les messages d'ifctester/xmlschema peuvent citer le chemin temporaire
        # du serveur : jamais renvoyé à l'utilisateur.
        reason = reason.replace(file_path, "le fichier")
        reason = re.sub(r"(?:/[\w.\- ]+){2,}", "<chemin>", reason)[:300]
        raise IdsFileError(f"Fichier IDS invalide : {reason}")
    if not document.specifications:
        raise IdsFileError("Fichier IDS sans aucune spécification")
    return document


def _cardinality(spec) -> str:
    if spec.maxOccurs == 0:
        return "prohibited"
    if spec.minOccurs != 0:
        return "required"
    return "optional"


def validate_lot(ids_document, ifc_file, floor_of: Optional[Callable] = None) -> dict:
    """Valide `ifc_file` contre `ids_document` et renvoie un résumé sérialisable.

    `floor_of(element) -> str` donne l'étage d'un élément (injecté par
    l'appelant, qui connaît déjà cette logique) ; sans lui, "Sans étage".

    Les spécifications dont la version IFC ne vise pas ce modèle sont
    ignorées (comportement normatif de l'IDS), pas comptées en échec.
    """
    ids_document.validate(ifc_file, should_filter_version=True)

    floor_cache = {}

    def floor(element) -> str:
        key = element.id()
        if key not in floor_cache:
            try:
                floor_cache[key] = str(floor_of(element)) if floor_of else "Sans étage"
            except Exception:
                floor_cache[key] = "Sans étage"
        return floor_cache[key]

    specifications = []
    groups = Counter()          # (spec_name, floor, ifc_class) -> nb d'éléments distincts
    details = []
    details_truncated = False
    applicable_total = 0
    failed_total = 0
    specs_checked = 0
    specs_failed = 0
    cardinality_failures = 0

    for spec in ids_document.specifications:
        cardinality = _cardinality(spec)
        base = {
            "name": spec.name,
            "identifier": getattr(spec, "identifier", None),
            "description": spec.description,
            "instructions": spec.instructions,
            "cardinality": cardinality,
        }
        if spec.status is None:
            specifications.append({
                **base, "status": "ignoree", "applicable": 0, "failed": 0, "passed": 0, "requirements": [],
            })
            continue

        applicable_entities = list(spec.applicable_entities)
        applicable = len(applicable_entities)

        if cardinality == "prohibited":
            failed_entities = applicable_entities
            failure_rows = [(e, "Interdit par cette spécification", "Élément interdit") for e in failed_entities]
        else:
            failed_entities = list(spec.failed_entities)
            failure_rows = []
            for requirement in spec.requirements:
                description = requirement.to_string("requirement", spec, requirement)
                for failure in requirement.failures:
                    failure_rows.append((failure["element"], str(failure["reason"]), description))

        for element in failed_entities:
            groups[(spec.name, floor(element), element.is_a())] += 1

        for element, reason, description in failure_rows:
            if len(details) >= MAX_DETAIL_ROWS_PER_LOT:
                details_truncated = True
                break
            details.append({
                "specification": spec.name,
                "requirement": description,
                "floor": floor(element),
                "ifc_class": element.is_a(),
                "guid": getattr(element, "GlobalId", None),
                "name": getattr(element, "Name", None),
                "reason": reason,
            })

        failed = len(failed_entities)
        status = "conforme" if spec.status else "non_conforme"
        specs_checked += 1
        applicable_total += applicable
        failed_total += failed
        if not spec.status:
            specs_failed += 1
            if cardinality == "required" and applicable == 0:
                cardinality_failures += 1

        specifications.append({
            **base,
            "status": status,
            "applicable": applicable,
            "failed": failed,
            "passed": applicable - failed,
            "requirements": [
                {
                    "description": requirement.to_string("requirement", spec, requirement),
                    "failed": len(requirement.failures),
                    "status": bool(requirement.status),
                }
                for requirement in spec.requirements
            ],
        })

    pass_rate = (applicable_total - failed_total) / applicable_total if applicable_total else None
    return {
        "error": None,
        "ifc_schema": ifc_file.schema,
        "specifications": specifications,
        "specs_checked": specs_checked,
        "specs_failed": specs_failed,
        "cardinality_failures": cardinality_failures,
        "applicable_total": applicable_total,
        "failed_total": failed_total,
        "pass_rate": pass_rate,
        "groups": [
            {"specification": spec_name, "floor": floor_name, "ifc_class": ifc_class, "count": count}
            for (spec_name, floor_name, ifc_class), count in groups.items()
        ],
        "details": details,
        "details_truncated": details_truncated,
    }


def error_result(message: str) -> dict:
    """Résultat factice pour un lot dont le contrôle IDS a échoué
    (l'analyse du reste n'est pas interrompue)."""
    return {
        "error": message, "ifc_schema": None, "specifications": [], "specs_checked": 0,
        "specs_failed": 0, "cardinality_failures": 0, "applicable_total": 0, "failed_total": 0,
        "pass_rate": None, "groups": [], "details": [], "details_truncated": False,
    }


def anomaly_label(spec_name: str) -> str:
    """Libellé stocké dans le champ « PSet » des anomalies pour un
    non-respect d'IDS (même structure que les anomalies PSet de l'Excel)."""
    return f"IDS : {spec_name}"


def public_summary(result: dict) -> dict:
    """Version allégée d'un résultat de lot pour la réponse JSON/le
    dashboard : sans le détail élément par élément (réservé à l'Excel)."""
    return {k: v for k, v in result.items() if k not in ("details", "groups")}


_STATUS_LABELS = {"conforme": "Conforme", "non_conforme": "Non conforme", "ignoree": "Ignorée (autre version IFC)"}
_CARDINALITY_LABELS = {"required": "Obligatoire", "optional": "Facultative", "prohibited": "Interdite"}


def create_ids_sheet(workbook, ids_title: str, ids_filename: str, results_by_lot: dict) -> None:
    """Onglet Excel "IDS" : synthèse par lot et spécification, puis détail
    des éléments non conformes."""
    sheet = workbook.create_sheet("IDS")

    title_font = Font(size=14, bold=True, color="FFFFFF")
    section_font = Font(size=12, bold=True)
    header_font = Font(size=11, bold=True)
    normal_font = Font(size=11)
    note_font = Font(size=10, italic=True, color="777777")
    title_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")
    section_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ko_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    grey_fill = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    n_cols = 8

    def banner(row, text, font, fill=None):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        cell = sheet.cell(row=row, column=1)
        cell.value = excel_safe(text)
        cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = center if font is title_font else left
        return row + 1

    def header(row, labels):
        for col, label in enumerate(labels, 1):
            cell = sheet.cell(row=row, column=col)
            cell.value = label
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        return row + 1

    def put(row, values, fills=None):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=col)
            cell.value = excel_safe(value)
            cell.font = normal_font
            cell.alignment = left if isinstance(value, str) else center
            cell.border = border
            if fills and fills.get(col):
                cell.fill = fills[col]

    r = banner(1, "CONTRÔLE IDS (Information Delivery Specification)", title_font, title_fill)
    sheet.row_dimensions[1].height = 28
    r += 1
    r = banner(r, f"Fichier IDS : {ids_filename}", normal_font)
    r = banner(r, f"Titre : {ids_title}", normal_font)
    r += 1

    r = banner(r, "1. Synthèse par lot et par spécification", section_font, section_fill)
    r = header(r, ["Lot", "Spécification", "Cardinalité", "Statut", "Éléments applicables",
                   "Non conformes", "Conformes", "Exigences non respectées"])
    for lot_name, result in results_by_lot.items():
        if result.get("error"):
            put(r, [lot_name, f"Contrôle IDS impossible : {result['error']}", "", "", "", "", "", ""],
                {2: ko_fill})
            r += 1
            continue
        for spec in result["specifications"]:
            failed_reqs = "; ".join(q["description"] for q in spec["requirements"] if not q["status"])
            fill = {"conforme": ok_fill, "non_conforme": ko_fill}.get(spec["status"], grey_fill)
            put(r, [
                lot_name, spec["name"], _CARDINALITY_LABELS.get(spec["cardinality"], spec["cardinality"]),
                _STATUS_LABELS.get(spec["status"], spec["status"]),
                spec["applicable"], spec["failed"], spec["passed"], failed_reqs,
            ], {4: fill})
            r += 1
    r += 1

    r = banner(r, "2. Détail des éléments non conformes", section_font, section_fill)
    r = header(r, ["Lot", "Spécification", "Exigence", "Étage", "Classe IFC", "GlobalId", "Nom", "Raison"])
    any_rows = False
    for lot_name, result in results_by_lot.items():
        for d in result.get("details", []):
            put(r, [lot_name, d["specification"], d["requirement"], d["floor"],
                    d["ifc_class"], d["guid"] or "", d["name"] or "", d["reason"]])
            r += 1
            any_rows = True
        if result.get("details_truncated"):
            r = banner(
                r, f"Lot « {lot_name} » : liste tronquée à {MAX_DETAIL_ROWS_PER_LOT} lignes "
                   f"(les compteurs de la synthèse restent exacts).", note_font,
            )
    if not any_rows:
        r = banner(r, "Aucun élément non conforme.", normal_font)

    widths = {"A": 22, "B": 38, "C": 46, "D": 22, "E": 20, "F": 26, "G": 24, "H": 52}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A5"
